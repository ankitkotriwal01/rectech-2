import streamlit as st
import io
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, RGBColor

# Create a new Word document
doc = Document()

# Default font settings
default_font_name = 'Calibri'
default_font_size = Pt(12)
default_paragraph_alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
default_font_color = RGBColor(0, 0, 0)

# Apply styles to the Normal style
normal_style = doc.styles['Normal']
normal_style.font.name = default_font_name
normal_style.font.size = default_font_size
normal_style.paragraph_format.alignment = default_paragraph_alignment
normal_style.font.color.rgb = default_font_color

# Streamlit app title and introduction
st.title("Cover Letter Generator")

# Get user input for personal and job details
full_name = st.text_input("Enter your full name:").strip()
address = st.text_input("Enter your address:").strip()
phone_number = st.text_input("Enter your phone number:").strip()
email_address = st.text_input("Enter your email address:").strip()
company_name = st.text_input("Enter the company name you're applying to:").strip()
hiring_manager = st.text_input("Enter the hiring manager's name (if known):").strip()
position = st.text_input("Enter the position you're applying for:").strip()
about_you = st.text_area("Tell us about yourself (your background, skills, and experiences):")
why_company = st.text_area("Why do you want to work at this company? What interests you about this position?")

# Start building the cover letter
doc.add_paragraph(f"{full_name}\n{address}\n{phone_number}\n{email_address}\n\n")
doc.add_paragraph().add_run(f"To {hiring_manager or 'Hiring Manager'},\n{company_name}\n\n").bold = True

# Opening paragraph
opening_paragraph = f"Dear {hiring_manager or 'Hiring Manager'},\n\nI am writing to express my interest in the {position} position at {company_name} as advertised. With my background in [Your Field/Industry] and extensive experience in [Relevant Experience], I am confident in my ability to contribute effectively to your team."
doc.add_paragraph(opening_paragraph)

# About you
doc.add_heading("About Me:", level=2)
doc.add_paragraph(about_you)

# Why the company
doc.add_heading("Why [Company Name]:", level=2)
doc.add_paragraph(why_company.replace('[Company Name]', company_name))


# Closing paragraph
closing_paragraph = "I am excited about the opportunity to bring my unique talents to [Company Name], a place known for [Something Noteworthy about the Company]. I look forward to the possibility of discussing this exciting opportunity with you. Thank you for considering my application. I am eager to bring my background in [Your Field/Industry] to your team and make a positive impact."
doc.add_paragraph(closing_paragraph.replace('[Company Name]', company_name))

# Signature
doc.add_paragraph("\nSincerely,\n\n" + full_name)

# Save the cover letter to a file
file_name = st.text_input("Enter the file name for saving the cover letter")

if file_name:
    file_path = f"{file_name}.docx"

    # Function to download the cover letter
    def download_cover_letter():
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output.getvalue()

    st.download_button(label="Download Cover Letter", data=download_cover_letter(), file_name=file_path, mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
